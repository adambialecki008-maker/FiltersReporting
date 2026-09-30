from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from filters_reporting.config import REPORTS_DIR

SUMMARY_HEADERS = [
    "Filter",
    "Min Δp",
    "Max Δp",
    "Średnia Δp",
    "Próg Δp [Pa]",
    "Liczba próbek",
    "Alarm aktywny",
    "Praca",
    "Aktywne alarmy %",
    "Praca %",
]

SAMPLE_HEADERS = [
    "ID",
    "Filtr",
    "ID Filtra",
    "Δp",
    "Alarm",
    "Praca",
    "Status",
    "Czas",
    "Czas wykresu",
]


STATUS_COLORS = {
    "Awaria": "FECACA",
    "Brudny – praca": "FED7AA",
    "Brudny – postój": "FED7AA",
    "Czysty – praca": "DCFCE7",
    "Czysty – postój": "DCFCE7",
}


def create_excel_report(
    stats,
    samples,
    chosen_day,
):
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Podsumowanie"

    samples_sheet = workbook.create_sheet("Próbki")

    create_summary_report(
        summary_sheet,
        stats,
    )

    row_ranges = create_samples_report(
        samples_sheet,
        samples,
        stats,
    )

    add_filter_charts(
        samples_sheet,
        row_ranges,
    )

    setup_summary_sheet(summary_sheet)

    setup_samples_sheet(samples_sheet)

    REPORTS_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    filepath = REPORTS_DIR / f"Raport filtry_{chosen_day}.xlsx"

    saved = save_excel_report(
        workbook,
        filepath,
    )

    if not saved:
        return False

    return filepath


def create_summary_report(
    worksheet,
    stats,
):
    worksheet.append(SUMMARY_HEADERS)

    for stat in stats:
        if stat.sample_count > 0:
            alarm_percentage = stat.alarm_count / stat.sample_count

            working_percentage = stat.status_count / stat.sample_count

        else:
            alarm_percentage = "Brak"
            working_percentage = "Brak"

        worksheet.append(
            [
                stat.filter_name,
                stat.min_delta_p,
                stat.max_delta_p,
                round(stat.avg_delta_p),
                stat.delta_p_threshold,
                stat.sample_count,
                stat.alarm_count,
                stat.status_count,
                alarm_percentage,
                working_percentage,
            ]
        )

    for row in range(
        2,
        worksheet.max_row + 1,
    ):
        alarm_cell = worksheet.cell(
            row=row,
            column=9,
        )

        work_cell = worksheet.cell(
            row=row,
            column=10,
        )

        if isinstance(
            alarm_cell.value,
            (int, float),
        ):
            alarm_cell.number_format = "0.0%"

        if isinstance(
            work_cell.value,
            (int, float),
        ):
            work_cell.number_format = "0.0%"


def create_samples_report(
    worksheet,
    samples,
    stats,
):
    worksheet.append(SAMPLE_HEADERS)

    threshold_by_filter = {stat.filter_name: stat.delta_p_threshold for stat in stats}

    last_working_delta_p = {}

    row_ranges = {}

    for sample in samples:
        threshold = threshold_by_filter.get(sample.filter_name)

        previous_working_delta_p = last_working_delta_p.get(sample.filter_id)

        business_status = get_report_sample_status(
            sample=sample,
            delta_p_threshold=(threshold),
            last_working_delta_p=(previous_working_delta_p),
        )

        worksheet.append(
            [
                sample.sample_id,
                sample.filter_name,
                sample.filter_id,
                sample.delta_p,
                boolean_text(sample.alarm_active),
                boolean_text(sample.status),
                business_status,
                sample.timestamp,
                sample.timestamp.strftime("%H:%M"),
            ]
        )

        current_row = worksheet.max_row

        if sample.filter_name not in row_ranges:
            row_ranges[sample.filter_name] = [
                current_row,
                current_row,
            ]

        else:
            row_ranges[sample.filter_name][1] = current_row

        status_cell = worksheet.cell(
            row=current_row,
            column=7,
        )

        status_color = STATUS_COLORS.get(business_status)

        if status_color:
            status_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=status_color,
            )

        worksheet.cell(
            row=current_row,
            column=8,
        ).number_format = "yyyy-mm-dd hh:mm:ss"

        if sample.status:
            last_working_delta_p[sample.filter_id] = sample.delta_p

    return row_ranges


def get_report_sample_status(
    sample,
    delta_p_threshold,
    last_working_delta_p,
):
    if sample.alarm_active:
        return "Awaria"

    if delta_p_threshold is None:
        if sample.status:
            return "Praca"

        return "Postój"

    if sample.status:
        if sample.delta_p >= delta_p_threshold:
            return "Brudny – praca"

        return "Czysty – praca"

    if last_working_delta_p is None:
        return "Czysty – postój"

    if last_working_delta_p >= delta_p_threshold:
        return "Brudny – postój"

    return "Czysty – postój"


def boolean_text(
    value,
):
    if value:
        return "TAK"

    return "NIE"


def add_filter_charts(
    worksheet,
    row_ranges,
):
    chart_number = 0

    for (
        filter_name,
        row_range,
    ) in row_ranges.items():
        start_row, end_row = row_range

        if end_row < start_row:
            continue

        chart = LineChart()

        chart.title = f"{filter_name} — ΔP [Pa]"

        chart.y_axis.title = "ΔP [Pa]"

        chart.x_axis.title = "Czas"

        chart.height = 7
        chart.width = 14

        data = Reference(
            worksheet,
            min_col=4,
            min_row=start_row,
            max_row=end_row,
        )

        categories = Reference(
            worksheet,
            min_col=9,
            min_row=start_row,
            max_row=end_row,
        )

        chart.add_data(
            data,
            titles_from_data=False,
        )

        series = chart.series[0]

        series.graphicalProperties.line.solidFill = "4F81BD"

        series.graphicalProperties.line.width = 38100

        chart.set_categories(categories)

        chart.legend = None

        anchor_row = 2 + chart_number * 15

        worksheet.add_chart(
            chart,
            f"K{anchor_row}",
        )

        chart_number += 1


def setup_summary_sheet(
    worksheet,
):
    style_header(worksheet)

    widths = {
        "A": 20,
        "B": 14,
        "C": 14,
        "D": 16,
        "E": 17,
        "F": 18,
        "G": 17,
        "H": 14,
        "I": 22,
        "J": 16,
    }

    for (
        column,
        width,
    ) in widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = worksheet.dimensions


def setup_samples_sheet(
    worksheet,
):
    style_header(worksheet)

    widths = {
        "A": 10,
        "B": 18,
        "C": 12,
        "D": 12,
        "E": 12,
        "F": 12,
        "G": 22,
        "H": 22,
        "I": 12,
    }

    for (
        column,
        width,
    ) in widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.column_dimensions["I"].hidden = True

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = f"A1:H{worksheet.max_row}"

    for row in worksheet.iter_rows(
        min_row=2,
        max_col=8,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
            )


def style_header(
    worksheet,
):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 22


def save_excel_report(
    workbook,
    filepath,
):
    workbook.save(filepath)

    return filepath
