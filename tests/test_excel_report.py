from datetime import datetime

from openpyxl import load_workbook

from filters_reporting.models.filter_report_models import (
    DailyFilterStats,
    DailyFilterReportSample,
)
from filters_reporting.reporting.excel_report import (
    boolean_text,
    create_excel_report,
    get_report_sample_status,
)


def make_sample(
    *,
    sample_id=1,
    filter_name="F1",
    filter_id=1,
    delta_p=1000,
    alarm_active=False,
    status=True,
):
    return DailyFilterReportSample(
        sample_id=sample_id,
        filter_name=filter_name,
        filter_id=filter_id,
        delta_p=delta_p,
        alarm_active=alarm_active,
        status=status,
        timestamp=datetime(
            2026,
            9,
            21,
            12,
            0,
        ),
    )


def make_stats(
    *,
    filter_name="F1",
    threshold=1500,
):
    return DailyFilterStats(
        filter_name=filter_name,
        min_delta_p=500,
        max_delta_p=2000,
        avg_delta_p=1250,
        sample_count=4,
        alarm_count=1,
        status_count=3,
        delta_p_threshold=threshold,
    )


def test_boolean_text_returns_tak_for_true():
    assert boolean_text(True) == "TAK"


def test_boolean_text_returns_nie_for_false():
    assert boolean_text(False) == "NIE"


def test_report_status_returns_error_when_alarm_active():
    sample = make_sample(
        alarm_active=True,
        status=True,
        delta_p=500,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=None,
    )

    assert result == "Awaria"


def test_report_status_returns_clean_working_below_threshold():
    sample = make_sample(
        status=True,
        delta_p=1499,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=None,
    )

    assert result == "Czysty – praca"


def test_report_status_returns_dirty_working_at_threshold():
    sample = make_sample(
        status=True,
        delta_p=1500,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=None,
    )

    assert result == "Brudny – praca"


def test_report_status_returns_clean_stopped_when_never_worked():
    sample = make_sample(
        status=False,
        delta_p=2000,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=None,
    )

    assert result == "Czysty – postój"


def test_report_status_returns_clean_stopped_from_last_working_delta_p():
    sample = make_sample(
        status=False,
        delta_p=2000,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=1000,
    )

    assert result == "Czysty – postój"


def test_report_status_returns_dirty_stopped_from_last_working_delta_p():
    sample = make_sample(
        status=False,
        delta_p=100,
    )

    result = get_report_sample_status(
        sample,
        delta_p_threshold=1500,
        last_working_delta_p=1600,
    )

    assert result == "Brudny – postój"


def test_create_excel_report_creates_file(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    stats = [
        make_stats(),
    ]

    samples = [
        make_sample(),
    ]

    filepath = create_excel_report(
        stats,
        samples,
        "2026-09-21",
    )

    assert filepath.exists()

    assert filepath.name == ("Raport filtry_2026-09-21.xlsx")


def test_create_excel_report_creates_expected_sheets(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    filepath = create_excel_report(
        [make_stats()],
        [make_sample()],
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    assert workbook.sheetnames == [
        "Podsumowanie",
        "Próbki",
    ]


def test_samples_sheet_has_expected_headers(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    filepath = create_excel_report(
        [make_stats()],
        [make_sample()],
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Próbki"]

    headers = [
        sheet.cell(
            row=1,
            column=column,
        ).value
        for column in range(1, 10)
    ]

    assert headers == [
        "ID",
        "Filtr",
        "ID Filtra",
        "Δp",
        "Alarm",
        "Praca",
        "Status",
        "Czas",
        "Czas wykresu",
    ]


def test_samples_sheet_contains_business_status(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    stats = [
        make_stats(threshold=1500),
    ]

    samples = [
        make_sample(
            sample_id=1,
            delta_p=1000,
            status=True,
        ),
        make_sample(
            sample_id=2,
            delta_p=1600,
            status=True,
        ),
        make_sample(
            sample_id=3,
            delta_p=100,
            status=False,
        ),
    ]

    filepath = create_excel_report(
        stats,
        samples,
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Próbki"]

    assert sheet["G2"].value == ("Czysty – praca")

    assert sheet["G3"].value == ("Brudny – praca")

    assert sheet["G4"].value == ("Brudny – postój")


def test_samples_sheet_has_separate_working_column(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    samples = [
        make_sample(
            sample_id=1,
            status=True,
        ),
        make_sample(
            sample_id=2,
            status=False,
        ),
    ]

    filepath = create_excel_report(
        [make_stats()],
        samples,
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Próbki"]

    assert sheet["F2"].value == "TAK"
    assert sheet["F3"].value == "NIE"


def test_summary_contains_filter_threshold(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    filepath = create_excel_report(
        [make_stats(threshold=1750)],
        [make_sample()],
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Podsumowanie"]

    assert sheet["E1"].value == ("Próg Δp [Pa]")

    assert sheet["E2"].value == 1750


def test_summary_percentages_are_calculated(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    filepath = create_excel_report(
        [make_stats()],
        [make_sample()],
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Podsumowanie"]

    assert sheet["I2"].value == 0.25
    assert sheet["J2"].value == 0.75


def test_samples_sheet_creates_one_chart_per_filter(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    stats = [
        make_stats(filter_name="F1"),
        make_stats(filter_name="F2"),
    ]

    samples = [
        make_sample(
            sample_id=1,
            filter_name="F1",
            filter_id=1,
        ),
        make_sample(
            sample_id=2,
            filter_name="F2",
            filter_id=2,
        ),
    ]

    filepath = create_excel_report(
        stats,
        samples,
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Próbki"]

    assert len(sheet._charts) == 2


def test_each_filter_chart_has_only_one_data_series(
    tmp_path,
    monkeypatch,
):
    monkeypatch.chdir(tmp_path)

    samples = [
        make_sample(
            sample_id=1,
            delta_p=1000,
        ),
        make_sample(
            sample_id=2,
            delta_p=1200,
        ),
        make_sample(
            sample_id=3,
            delta_p=1400,
        ),
    ]

    filepath = create_excel_report(
        [make_stats()],
        samples,
        "2026-09-21",
    )

    workbook = load_workbook(filepath)

    sheet = workbook["Próbki"]

    assert len(sheet._charts) == 1

    chart = sheet._charts[0]

    assert len(chart.series) == 1
